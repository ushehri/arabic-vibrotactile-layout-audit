"""S5.1 figures + Excel sheet '5.1'."""
import sys; from pathlib import Path; sys.path.insert(0, str(Path(__file__).resolve().parents[1]))  # code/ on the path, for lmp_paths
from lmp_paths import arabic_font, data, results_dir, figures_dir, workdir
import json, pandas as pd, numpy as np, matplotlib; matplotlib.use('Agg'); import matplotlib.pyplot as plt
from matplotlib import font_manager as fm; import arabic_reshaper; from bidi.algorithm import get_display
import openpyxl
fp=arabic_font()
if fp: fm.fontManager.addfont(fp)
AF=fm.FontProperties(family=['Noto Sans Arabic','DejaVu Sans'])
A=lambda s:get_display(arabic_reshaper.reshape(s))
R=results_dir('s5.1'); F=figures_dir()
J=json.load(open(R+'s51_results.json')); T=['ت','ر','ض','ي','ك','ق','خ','غ']; IPA=['t','r','dˤ','j','k','q','x','ɣ']
L=J['layout']['tip']; EV=J['evidence']
# Fig 1: layout position vs each evidence ordering
names={'KAPD_label':'KAPD place label','KAPD_EPG':'KAPD EPG mean contact position','Modern_lit':'Modern articulatory literature','Classical':'Classical makharij'}
fig,axs=plt.subplots(1,4,figsize=(15,4))
for a,k in zip(axs,names):
    y=np.array(EV[k],float); rng=np.random.default_rng(1)
    for i,l in enumerate(T):
        a.scatter(L[i],y[i],s=30,c='C0'); a.annotate(A(l)+f' /{IPA[i]}/',(L[i],y[i]),xytext=(4,3+8*((i%3)-1)),textcoords='offset points',fontproperties=AF,fontsize=9)
    r=J['results'][f'tip|{k}']; a.set_title(f"{names[k]}\nρ={r['spearman']:.2f}, τb={r['kendall_tau_b']:.2f}, discordant pairs={r['discordant']}",fontsize=9)
    a.set_xlabel('Sleeve position (cm from wrist)'); a.set_ylabel('front → back (evidence code)')
plt.tight_layout(); plt.savefig(F+'s51_fig1_layout_vs_evidence.png',dpi=150); plt.close()
# Fig 2: KAPD EPG per speaker
e=pd.read_csv(data('kapd/kapd_epg_summary.csv')); C={'TS':'ت','RS':'ر','DB':'ض','YS':'ي','KS':'ك','QS':'ق','XS':'خ','GS':'غ'}
e=e[e.code.isin(C)]; e['L']=e.code.map(C)
fig,a=plt.subplots(figsize=(8,4.2))
for s,m in zip('AFGM','os^D'):
    g=e[e.subj==s].groupby('L').mean_rel_y.mean().reindex(T)
    a.plot(range(8),g.values,m+'-',alpha=.6,label=f'speaker {s}')
g=e.groupby('L').mean_rel_y.mean().reindex(T); a.plot(range(8),g.values,'k-',lw=2.5,label='pooled')
a.set_xticks(range(8)); a.set_xticklabels([A(l)+f'\n/{p}/\n{c} cm' for l,p,c in zip(T,IPA,L)],fontproperties=AF)
a.set_ylabel('EPG mean contact position\n(0 = front, 1 = rearmost electrode)'); a.legend(fontsize=8); a.set_title('KAPD EPG (medial singletons, mean of /a i u/) in sleeve order',fontsize=10)
plt.tight_layout(); plt.savefig(F+'s51_fig2_kapd_epg_by_speaker.png',dpi=150); plt.close()
# Excel
import os; X=os.path.join(str(workdir('tables')),'letter-mapping-tables.xlsx')
ag=pd.read_csv(R+'s51_agreement.csv'); pl=pd.read_csv(R+'s51_per_letter_ranks.csv'); ep=pd.read_csv(R+'s51_kapd_epg_by_letter.csv')
mm=pd.read_csv(R+'s51_mismatch_table.csv')
wb=openpyxl.load_workbook(X) if os.path.exists(X) else openpyxl.Workbook()
if '5.1' in wb.sheetnames: del wb['5.1']
ws=wb.create_sheet('5.1')
def put(title,df):
    ws.append([title]); ws.append(list(df.columns))
    for r in df.itertuples(index=False): ws.append(list(r))
    ws.append([])
put('Table 5.1-1 Per-letter comparison and mismatch findings',mm)
put('Table 5.1-2 Rank agreement, layout vs evidence (exact permutation over 8! orders)',ag)
put('Table 5.1-2b Extra modern data sources (status)',pd.read_csv(R+'s51_agreement_extra.csv'))
import json as _j
u=_j.load(open(R+'s51_uti_emphasis.json'))
put('Table 5.1-4 UTI emphatic minus plain (X_z; negative = backer), Lebanese Arabic, Sakr data via Coretta & Sakr 2025',pd.DataFrame([dict(pair_region=k,**v) for k,v in u.items()]))
put('Table 5.1-3 KAPD EPG summary by letter (4 speakers x 3 vowels)',ep)
put('Per-letter ranks by evidence line',pl)
wb.save(X); print('ok')
# Fig 3: UTI emphasis (d vs dˤ) per region, per speaker
c=pd.read_csv(R+'s51_uti_emphasis_cells.csv'); c=c[(c.pair=='d')&(c.region!='Vallecula')]
order=['Root','Body','Dorsum','Blade','Tip']
fig,a=plt.subplots(figsize=(7,3.8))
for k,s in enumerate(sorted(c.participant.unique())):
    g=c[c.participant==s].groupby('region').dX.mean().reindex(order); a.plot(range(5),g.values,'o-',alpha=.6,label=s)
g=c.groupby('region').dX.mean().reindex(order); a.plot(range(5),g.values,'k-',lw=2.5,label='mean')
a.axhline(0,color='grey',lw=.8); a.set_xticks(range(5)); a.set_xticklabels(['root','body','dorsum','blade','tip'])
a.set_ylabel('emphatic − plain X (SD)\nnegative = more retracted'); a.set_title('/dˤ/ vs /d/, Lebanese Arabic UTI (5 speakers × 5 vowels)',fontsize=10); a.legend(fontsize=7,ncol=3)
plt.tight_layout(); plt.savefig(F+'s51_fig3_uti_emphasis.png',dpi=150); plt.close()
